"""Lectura de adjuntos para Brian (funciones puras, sin modelos ni superficie RPC).

* ``tipo_de_archivo(crudo, nombre, mimetype)`` decide por la FIRMA del archivo (y luego por la
  extensión): un navegador puede etiquetar un CSV como ``application/vnd.ms-excel`` y un .xls
  binario jamás debe leerse como texto.
* ``leer_excel_xlsx`` / ``leer_excel_xls`` / ``leer_docx`` devuelven texto compacto (tablas en
  markdown por hoja) dentro de un presupuesto de caracteres.
* ``normalizar_imagen`` deja una foto lista para el modelo de visión: lado mayor ≤ 1568 px,
  JPEG (o PNG si tiene transparencia), sin tocar el adjunto original.

Cualquier excepción la atrapa quien llama (``brian.conversacion._leer_adjunto``) y la persona
ve un mensaje amable.
"""
import datetime
import io
import re
import zipfile
from xml.etree import ElementTree

from PIL import Image

from odoo.tools.image import IMAGE_MAX_RESOLUTION, image_fix_orientation

FIRMA_ZIP = b'PK\x03\x04'
FIRMA_OLE = b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1'
FIRMA_PDF = b'%PDF-'

MAX_COLUMNAS = 30
MAX_CELDA = 80
FILAS_ENCABEZADO = 10
MAX_LADO_IMAGEN = 1568
CALIDAD_JPEG = 85


def _extension(nombre):
    nombre = (nombre or '').lower()
    return nombre.rsplit('.', 1)[-1] if '.' in nombre else ''


def parece_texto(crudo):
    """Sin bytes NUL al inicio: probablemente texto (CSV, TXT…)."""
    return b'\x00' not in crudo[:4096]


def tipo_de_archivo(crudo, nombre, mimetype):
    """'xlsx' | 'xls' | 'docx' | 'pdf' | 'protegido' | 'danado' | 'zip' | 'ole' | None.

    ``None`` = sin firma conocida: se decide por mimetype/extensión (texto o no legible).
    """
    ext = _extension(nombre)
    if crudo.startswith(FIRMA_ZIP):
        try:
            nombres = set(zipfile.ZipFile(io.BytesIO(crudo)).namelist())
        except zipfile.BadZipFile:
            # Dice ser un .xlsx/.docx pero el zip está roto.
            return 'danado'
        if 'xl/workbook.xml' in nombres:
            return 'xlsx'
        if 'word/document.xml' in nombres:
            return 'docx'
        return 'zip'
    if crudo.startswith(FIRMA_OLE):
        if ext in ('xlsx', 'xlsm', 'docx'):
            # Un Office moderno cifrado con contraseña es un contenedor OLE.
            return 'protegido'
        if ext == 'xls' or 'excel' in (mimetype or ''):
            return 'xls'
        return 'ole'
    if crudo.startswith(FIRMA_PDF):
        return 'pdf'
    if ext in ('xlsx', 'xlsm', 'docx') or (ext == 'xls' and not parece_texto(crudo)):
        return 'danado'
    return None


def _numero(valor):
    """Números sin ruido de coma flotante: 418.65999999999997 → 418.66, 12.0 → 12."""
    redondeado = round(valor, 2)
    if redondeado == int(redondeado):
        return str(int(redondeado))
    return f'{redondeado:.2f}'.rstrip('0').rstrip('.')


def _celda(valor):
    if valor is None:
        return ''
    if isinstance(valor, bool):
        return 'Sí' if valor else 'No'
    if isinstance(valor, float):
        valor = _numero(valor)
    elif isinstance(valor, datetime.datetime):
        valor = valor.strftime('%d/%m/%Y %H:%M') if (valor.hour or valor.minute) else valor.strftime('%d/%m/%Y')
    elif isinstance(valor, datetime.date):
        valor = valor.strftime('%d/%m/%Y')
    texto = ' '.join(str(valor).split()).replace('|', '/')
    return texto if len(texto) <= MAX_CELDA else texto[:MAX_CELDA - 1] + '…'


def _celdas(fila):
    celdas = [_celda(v) for v in list(fila)[:MAX_COLUMNAS]]
    while celdas and not celdas[-1]:
        celdas.pop()
    return celdas


def _hoja(nombre, filas, total_filas, total_columnas, presupuesto, imagenes=0):
    """Una hoja como texto compacto, recortada al ``presupuesto`` de caracteres.

    ``filas`` es un iterable de listas de valores. El encabezado es la fila con más celdas llenas
    entre las primeras ``FILAS_ENCABEZADO`` (los catálogos suelen traer título y notas arriba):
    lo de antes se muestra como contexto y lo de después como tabla markdown.
    """
    lineas = [f'### Hoja «{nombre}» ({total_filas} filas × {total_columnas} columnas)']
    if imagenes:
        lineas.append(('(1 imagen incrustada' if imagenes == 1 else f'({imagenes} imágenes incrustadas')
                      + ': no las ves, pero proponer_importacion pone cada una como foto del producto de su fila.)')
    filas = iter(filas)
    primeras = []
    for fila in filas:
        celdas = _celdas(fila)
        if celdas:
            primeras.append(celdas)
        if len(primeras) >= FILAS_ENCABEZADO:
            break
    if not primeras:
        lineas.append('(Hoja vacía)')
        return '\n'.join(lineas)
    llenas = [sum(1 for c in celdas if c) for celdas in primeras]
    indice = llenas.index(max(llenas))
    for celdas in primeras[:indice]:
        lineas.append(' · '.join(c for c in celdas if c))
    encabezado = primeras[indice]
    lineas.append('| ' + ' | '.join(encabezado) + ' |')
    lineas.append('|' + '---|' * len(encabezado))
    usado = sum(len(linea) + 1 for linea in lineas)
    escritas = vistas = 0

    def resto():
        yield from primeras[indice + 1:]
        for fila in filas:
            yield _celdas(fila)

    for celdas in resto():
        if not celdas:
            continue
        vistas += 1
        linea = '| ' + ' | '.join(celdas) + ' |'
        if escritas < vistas - 1 or usado + len(linea) + 1 > presupuesto:
            continue
        lineas.append(linea)
        usado += len(linea) + 1
        escritas += 1
    if escritas < vistas:
        lineas.append(f'… {vistas - escritas} filas más (no caben aquí; pídeme un resumen o filtra el archivo).')
    if total_columnas > MAX_COLUMNAS:
        lineas.append(f'(Solo se muestran las primeras {MAX_COLUMNAS} columnas.)')
    return '\n'.join(lineas)


def _unir(textos_por_hoja, presupuesto):
    # Cada hoja ya respetó su cuota; esto solo es la red de seguridad.
    return '\n\n'.join(textos_por_hoja)[: presupuesto + 2000]


def _imagenes_por_hoja_xlsx(crudo):
    """{nombre de hoja: cuántas imágenes}. Solo carga el libro completo si trae imágenes."""
    nombres = zipfile.ZipFile(io.BytesIO(crudo)).namelist()
    if not any(n.startswith('xl/media/') for n in nombres):
        return {}
    import openpyxl
    libro = openpyxl.load_workbook(io.BytesIO(crudo), data_only=True)
    try:
        return {hoja.title: len(getattr(hoja, '_images', []) or []) for hoja in libro.worksheets}
    finally:
        libro.close()


def _con_formulas(valores, formulas):
    """Valor guardado de cada celda; si una fórmula no tiene valor guardado, su texto marcado."""
    for fila_valores, fila_formulas in zip(valores, formulas, strict=False):
        yield [v if v is not None or not (isinstance(f, str) and f.startswith('='))
               else f'{f} (fórmula sin valor guardado)'
               for v, f in zip(fila_valores, fila_formulas, strict=False)]


def leer_excel_xlsx(crudo, presupuesto):
    import openpyxl
    imagenes = _imagenes_por_hoja_xlsx(crudo)
    libro = openpyxl.load_workbook(io.BytesIO(crudo), read_only=True, data_only=True)
    libro_formulas = openpyxl.load_workbook(io.BytesIO(crudo), read_only=True, data_only=False)
    try:
        hojas = libro.worksheets
        cuota = max(presupuesto // max(len(hojas), 1), 1500)
        textos = [f'Libro de Excel con {len(hojas)} hoja(s).']
        for hoja, hoja_formulas in zip(hojas, libro_formulas.worksheets, strict=True):
            filas = _con_formulas(hoja.iter_rows(values_only=True), hoja_formulas.iter_rows(values_only=True))
            textos.append(_hoja(hoja.title, filas, hoja.max_row or 0, hoja.max_column or 0, cuota,
                                imagenes.get(hoja.title, 0)))
    finally:
        libro.close()
        libro_formulas.close()
    return _unir(textos, presupuesto)


def leer_excel_xls(crudo, presupuesto):
    import xlrd
    libro = xlrd.open_workbook(file_contents=crudo, on_demand=True)
    try:
        hojas = [libro.sheet_by_index(i) for i in range(libro.nsheets)]
        cuota = max(presupuesto // max(len(hojas), 1), 1500)
        textos = [f'Libro de Excel (.xls) con {len(hojas)} hoja(s).']
        for hoja in hojas:
            def filas(h=hoja):
                for r in range(h.nrows):
                    valores = []
                    for celda in h.row(r):
                        if celda.ctype == xlrd.XL_CELL_DATE:
                            try:
                                valores.append(xlrd.xldate.xldate_as_datetime(celda.value, libro.datemode))
                                continue
                            except (ValueError, OverflowError, xlrd.xldate.XLDateError):
                                pass
                        valores.append(celda.value if celda.ctype != xlrd.XL_CELL_EMPTY else None)
                    yield valores
            textos.append(_hoja(hoja.name, filas(), hoja.nrows, hoja.ncols, cuota))
    finally:
        libro.release_resources()
    return _unir(textos, presupuesto)


_W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'


def leer_docx(crudo, presupuesto):
    """Texto de un .docx (párrafos y tablas) leyendo el XML; sin dependencias extra."""
    xml = zipfile.ZipFile(io.BytesIO(crudo)).read('word/document.xml')
    raiz = ElementTree.fromstring(xml)
    parrafos = []
    for parrafo in raiz.iter(f'{_W}p'):
        texto = ''.join(t.text or '' for t in parrafo.iter(f'{_W}t'))
        if texto.strip():
            parrafos.append(texto.strip())
    texto = '\n'.join(parrafos)
    return re.sub(r'\n{3,}', '\n\n', texto)[:presupuesto + 2000]


def normalizar_imagen(crudo, max_bytes):
    """(mimetype, bytes) listos para el modelo, o (None, motivo) si no se puede.

    No toca el original: si ya es pequeño y en formato común, se devuelve tal cual.
    """
    try:
        imagen = Image.open(io.BytesIO(crudo))
        formato = (imagen.format or '').upper()
        ancho, alto = imagen.size
    except Exception:  # noqa: BLE001 — HEIC u otro formato que PIL no abre
        return None, 'formato que no puedo abrir (si es HEIC, mándala como JPG o PNG)'
    if ancho * alto > IMAGE_MAX_RESOLUTION:
        return None, 'tiene demasiados píxeles'
    comunes = {'JPEG': 'image/jpeg', 'PNG': 'image/png', 'GIF': 'image/gif', 'WEBP': 'image/webp'}
    if formato in comunes and max(ancho, alto) <= MAX_LADO_IMAGEN and len(crudo) <= max_bytes:
        return comunes[formato], crudo
    imagen = image_fix_orientation(imagen)
    imagen.thumbnail((MAX_LADO_IMAGEN, MAX_LADO_IMAGEN), Image.Resampling.LANCZOS)
    transparente = imagen.mode in ('RGBA', 'LA') or (imagen.mode == 'P' and 'transparency' in imagen.info)
    intentos = [('PNG', 'image/png')] if transparente else []
    intentos.append(('JPEG', 'image/jpeg'))
    for salida, mimetype in intentos:
        lienzo = imagen
        if salida == 'JPEG' and lienzo.mode != 'RGB':
            lienzo = lienzo.convert('RGB')
        bufer = io.BytesIO()
        opciones = {'optimize': True}
        if salida == 'JPEG':
            opciones['quality'] = CALIDAD_JPEG
        lienzo.save(bufer, format=salida, **opciones)
        datos = bufer.getvalue()
        if len(datos) <= max_bytes:
            return mimetype, datos
    return None, 'es demasiado grande'
