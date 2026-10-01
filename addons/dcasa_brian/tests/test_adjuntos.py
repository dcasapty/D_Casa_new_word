"""Lectura de adjuntos: Excel (.xlsx/.xls), Word, CSV mal etiquetado y fotos grandes."""
import base64
import io
import os
import zipfile

from PIL import Image

from odoo.addons.dcasa_brian.models import conversacion as modulo_conversacion
from odoo.addons.dcasa_brian.models import lector_adjuntos
from odoo.tests import tagged

from .common import BrianCase

MIME_XLSX = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
AMABLE = 'No pude leer el archivo; puede estar dañado o protegido.'


def _png(ancho, alto, modo='RGB', ruido=False):
    if ruido:
        imagen = Image.frombytes(modo, (ancho, alto), os.urandom(ancho * alto * len(modo)))
    else:
        imagen = Image.new(modo, (ancho, alto), (19, 64, 177, 128) if modo == 'RGBA' else (19, 64, 177))
    bufer = io.BytesIO()
    imagen.save(bufer, format='PNG')
    return bufer.getvalue()


def _catalogo_xlsx(filas_extra=0):
    """Mismo «molde» que el catálogo real de un proveedor: título y notas arriba, encabezado en
    la fila 4, precios con ruido de coma flotante, una fórmula sin valor guardado y una foto."""
    import openpyxl
    from openpyxl.drawing.image import Image as ImagenXL
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = 'Catálogo LT-01'
    hoja.append(['Catálogo de productos – Pedido LT-01'])
    hoja.append(['Proveedor: Fábrica de Prueba · Fecha PI: 09/08/2026'])
    hoja.append(['Buscar (nombre o código):'])
    hoja.append(['Imagen', 'Código', 'Descripción', 'Precio cama sola', 'Combo'])
    hoja.append([None, '888K', 'Cama tapizada King – beige', 418.65999999999997, '=D5*2'])
    hoja.append([None, '888Q', 'Cama tapizada Queen', 278.1893, 556.38])
    for i in range(filas_extra):
        hoja.append([None, f'X{i:04d}', 'Relleno de prueba con una descripción larguita', 10.5, 21])
    foto = io.BytesIO(_png(20, 20))
    hoja.add_image(ImagenXL(foto), 'A5')
    notas = libro.create_sheet('Notas')
    notas.append(['Nota'])
    notas.append(['Precios sin ITBMS'])
    bufer = io.BytesIO()
    libro.save(bufer)
    return bufer.getvalue()


@tagged('post_install', '-at_install')
class TestAdjuntos(BrianCase):

    def leer(self, crudo, nombre, mimetype):
        adjunto = self.env['ir.attachment'].create({'name': nombre, 'raw': crudo, 'mimetype': mimetype})
        return self.env['brian.conversacion']._leer_adjunto(adjunto, mimetype)

    def test_xlsx_con_dos_hojas_titulo_e_imagenes(self):
        texto = self.leer(_catalogo_xlsx(), 'Catalogo_LT-01.xlsx', MIME_XLSX)
        self.assertIn('2 hoja(s)', texto)
        self.assertIn('### Hoja «Catálogo LT-01» (6 filas × 5 columnas)', texto)
        self.assertIn('1 imagen incrustada', texto)
        # Las filas de título quedan como contexto encima de la tabla.
        self.assertIn('Catálogo de productos – Pedido LT-01', texto)
        self.assertIn('| Imagen | Código | Descripción | Precio cama sola | Combo |', texto)
        self.assertLess(texto.index('Pedido LT-01'), texto.index('| Imagen |'))
        self.assertIn('| 888K | Cama tapizada King – beige | 418.66 | =D5*2 (fórmula sin valor guardado) |', texto)
        self.assertIn('| 888Q | Cama tapizada Queen | 278.19 | 556.38 |', texto)
        self.assertNotIn('418.659', texto)
        self.assertIn('### Hoja «Notas»', texto)
        self.assertIn('Precios sin ITBMS', texto)

    def test_xlsx_grande_se_recorta_con_aviso(self):
        texto = self.leer(_catalogo_xlsx(filas_extra=3000), 'grande.xlsx', MIME_XLSX)
        self.assertLessEqual(len(texto), modulo_conversacion.MAX_TEXTO_ADJUNTO + 200)
        self.assertRegex(texto, r'… \d+ filas más')
        self.assertIn('| Imagen | Código |', texto)
        self.assertIn('### Hoja «Notas»', texto, 'Cada hoja tiene su cuota')

    def test_xls_binario(self):
        import xlwt
        libro = xlwt.Workbook()
        hoja = libro.add_sheet('Inventario')
        for fila, valores in enumerate([('Código', 'Existencia'), ('SOF-01', 3), ('MES-02', 12.25)]):
            for columna, valor in enumerate(valores):
                hoja.write(fila, columna, valor)
        bufer = io.BytesIO()
        libro.save(bufer)
        crudo = bufer.getvalue()
        self.assertEqual(lector_adjuntos.tipo_de_archivo(crudo, 'inv.xls', 'application/vnd.ms-excel'), 'xls')
        texto = self.leer(crudo, 'inv.xls', 'application/vnd.ms-excel')
        self.assertIn('### Hoja «Inventario» (3 filas × 2 columnas)', texto)
        self.assertIn('| Código | Existencia |', texto)
        self.assertIn('| SOF-01 | 3 |', texto)
        self.assertIn('| MES-02 | 12.25 |', texto)

    def test_csv_etiquetado_como_excel_es_texto(self):
        crudo = b'codigo,precio\nSOF-01,10\n'
        self.assertIsNone(lector_adjuntos.tipo_de_archivo(crudo, 'export', 'application/vnd.ms-excel'))
        self.assertIn('SOF-01,10', self.leer(crudo, 'export', 'application/vnd.ms-excel'))
        self.assertIn('SOF-01,10', self.leer(crudo, 'export.csv', 'application/vnd.ms-excel'))

    def test_excel_danado_o_protegido_da_mensaje_amable(self):
        roto = lector_adjuntos.FIRMA_ZIP + os.urandom(200)
        self.assertIn(AMABLE, self.leer(roto, 'roto.xlsx', MIME_XLSX))
        self.assertIn(AMABLE, self.leer(os.urandom(300).replace(b'PK', b'pk'), 'basura.xlsx', MIME_XLSX))
        protegido = lector_adjuntos.FIRMA_OLE + os.urandom(500)
        self.assertIn(AMABLE, self.leer(protegido, 'clave.xlsx', MIME_XLSX))
        # Un .xls binario que no se puede abrir tampoco se «lee» como texto basura.
        self.assertIn(AMABLE, self.leer(protegido, 'viejo.xls', 'application/vnd.ms-excel'))

    def test_docx(self):
        documento = (
            '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>'
            '<w:p><w:r><w:t>Cotización para </w:t></w:r><w:r><w:t>la señora Ana</w:t></w:r></w:p>'
            '<w:tbl><w:tr><w:tc><w:p><w:r><w:t>Sofá en L</w:t></w:r></w:p></w:tc></w:tr></w:tbl>'
            '</w:body></w:document>')
        bufer = io.BytesIO()
        with zipfile.ZipFile(bufer, 'w') as archivo:
            archivo.writestr('word/document.xml', documento)
        texto = self.leer(bufer.getvalue(), 'cotizacion.docx',
                          'application/vnd.openxmlformats-officedocument.wordprocessingml.document')
        self.assertIn('Cotización para la señora Ana', texto)
        self.assertIn('Sofá en L', texto)

    # -- Fotos -------------------------------------------------------------------------

    def test_foto_grande_se_ajusta_sin_tocar_el_original(self):
        crudo = _png(2400, 1800, ruido=True)
        self.assertGreater(len(crudo), modulo_conversacion.MAX_IMAGEN)
        adjunto = self.env['ir.attachment'].create({'name': 'foto.png', 'raw': crudo, 'mimetype': 'image/png'})
        # (Odoo ya limita a 1920 px al guardar; lo que se compara es el adjunto tal como quedó.)
        original = adjunto.raw
        self.assertGreater(len(original), modulo_conversacion.MAX_IMAGEN)
        tipo, datos = modulo_conversacion._imagen_para_modelo(adjunto)
        self.assertEqual(tipo, 'image/jpeg')
        binario = base64.b64decode(datos)
        self.assertLessEqual(len(binario), modulo_conversacion.MAX_IMAGEN)
        self.assertEqual(max(Image.open(io.BytesIO(binario)).size), lector_adjuntos.MAX_LADO_IMAGEN)
        self.assertEqual(adjunto.raw, original, 'El adjunto original no cambia')

    def test_foto_con_transparencia_sigue_en_png(self):
        tipo, datos = lector_adjuntos.normalizar_imagen(_png(3000, 2000, 'RGBA'), modulo_conversacion.MAX_IMAGEN)
        self.assertEqual(tipo, 'image/png')
        imagen = Image.open(io.BytesIO(datos))
        self.assertEqual(max(imagen.size), lector_adjuntos.MAX_LADO_IMAGEN)
        self.assertIn(imagen.mode, ('RGBA', 'P'))

    def test_foto_pequena_queda_igual(self):
        crudo = _png(40, 30)
        self.assertEqual(lector_adjuntos.normalizar_imagen(crudo, modulo_conversacion.MAX_IMAGEN),
                         ('image/png', crudo))

    def test_heic_da_aviso(self):
        tipo, motivo = lector_adjuntos.normalizar_imagen(b'\x00\x00\x00\x18ftypheic' + os.urandom(64),
                                                         modulo_conversacion.MAX_IMAGEN)
        self.assertIsNone(tipo)
        self.assertIn('HEIC', motivo)
