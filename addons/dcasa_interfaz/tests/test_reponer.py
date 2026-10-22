"""«Reponer» (Compras › Reponer): ritmo de venta, alcanza para, tiempo de llegada, cuánto pedir y «Pedir ya»; el
buscador de proveedores, poner el tiempo de llegada, el Excel y los permisos. En el navegador, a 1440 y a 390 px.

Las pruebas del cálculo usan una ventana que termina el 31/05/2024 para no mezclarse con las ventas de otras
pruebas; las de la pantalla, el Excel y el navegador miran hasta hoy con productos y proveedores propios.
"""
import io
from datetime import date, timedelta

import openpyxl

from odoo import Command
from odoo.exceptions import AccessError, UserError
from odoo.tests import HttpCase, tagged
from odoo.tools import mute_logger

from .test_facturado import FacturadoCommon
from .test_inicio_por_rol import ESPERAR_JS, VentasMixin

HASTA = date(2024, 5, 31)          # la ventana de 5 meses va del 01/01/2024 al 31/05/2024


class ReponerDatos:
    """Proveedores y productos de prueba (nombres inventados, evidentes)."""

    @classmethod
    def crear_datos(cls):
        Socio = cls.env['res.partner']
        cls.prov_a = Socio.create({'name': 'PROVEEDOR DE PRUEBA Ásia', 'phone': '+507 6111-2233',
                                   'vat': '155000111-2-2020', 'supplier_rank': 1})
        cls.prov_b = Socio.create({'name': 'PROVEEDOR DE PRUEBA Bambú', 'supplier_rank': 1})
        cls.cliente_rep = Socio.create({'name': 'CLIENTE DE PRUEBA reponer'})
        Producto = cls.env['product.product']

        def producto(nombre, codigo, *precios):
            return Producto.create({
                'name': nombre, 'default_code': codigo, 'is_storable': True, 'list_price': 100.0,
                'taxes_id': [Command.clear()],
                'seller_ids': [Command.create(precio) for precio in precios]})
        cls.sofa = producto('Sofá reponer', 'REP-SOFA',
                            {'partner_id': cls.prov_a.id, 'price': 200.0, 'delay': 90},
                            {'partner_id': cls.prov_b.id, 'price': 190.0, 'sequence': 99})   # B no es el principal
        cls.mesa = producto('Mesa reponer', 'REP-MESA', {'partner_id': cls.prov_b.id, 'price': 70.0})  # 1 día: falta
        cls.silla = producto('Silla reponer', 'REP-SILLA')                                              # sin proveedor
        cls.lampara = producto('Lámpara reponer', 'REP-LAMP',
                               {'partner_id': cls.prov_a.id, 'price': 20.0, 'delay': 90})
        cls.bodega = cls.env['stock.warehouse'].search(
            [('company_id', '=', cls.env.company.id)], limit=1).lot_stock_id

    @classmethod
    def venta_facturada(cls, dia, producto, cantidad):
        factura = cls.env['account.move'].create({
            'move_type': 'out_invoice', 'partner_id': cls.cliente_rep.id, 'invoice_date': dia,
            'invoice_line_ids': [Command.create({'product_id': producto.id, 'quantity': cantidad,
                                                 'price_unit': 100.0, 'tax_ids': [Command.clear()]})]})
        factura.action_post()
        return factura

    @classmethod
    def poner_existencia(cls, producto, cantidad):
        cls.env['stock.quant']._update_available_quantity(producto, cls.bodega, cantidad)

    @classmethod
    def compra_al_proveedor(cls, proveedor, producto, cantidad, confirmar=True):
        compra = cls.env['purchase.order'].create({'partner_id': proveedor.id, 'order_line': [
            Command.create({'product_id': producto.id, 'product_qty': cantidad, 'price_unit': 10.0})]})
        if confirmar:
            compra.button_confirm()
        return compra

    @classmethod
    def precios_de(cls, proveedor):
        return cls.env['product.supplierinfo'].search([('partner_id', '=', proveedor.id)])


@tagged('post_install', '-at_install')
class TestReponer(ReponerDatos, FacturadoCommon):

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.crear_datos()

    def calcular(self, meses=5, reserva=0, proveedor=None, usuario=None):
        return self.env['dcasa.reponer'].with_user(usuario or self.gerencia)._calcular(meses, reserva, proveedor,
                                                                                       hasta=HASTA)

    def fila(self, filas, producto):
        return next(fila for fila in filas if fila['producto_id'] == producto.id)

    # --- El cálculo ------------------------------------------------------------------------------------------------

    def test_ritmo_alcanza_y_pedir_con_el_tiempo_de_llegada_y_lo_que_viene(self):
        self.venta_facturada(date(2024, 2, 10), self.sofa, 6)
        self.venta_facturada(date(2024, 5, 20), self.sofa, 4)                      # 10 en 5 meses → 2 al mes
        self.poner_existencia(self.sofa, 3)
        self.compra_al_proveedor(self.prov_a, self.sofa, 5)                        # viene en camino
        self.compra_al_proveedor(self.prov_a, self.sofa, 7, confirmar=False)       # borrador: no cuenta
        sofa = self.fila(self.calcular(), self.sofa)
        self.assertEqual(
            {clave: sofa[clave] for clave in ('vendido', 'ritmo', 'existencia', 'pedido', 'alcanza_meses',
                                              'llegada_dias', 'llegada', 'pedir', 'pedir_ya', 'estado')},
            {'vendido': 10.0, 'ritmo': 2.0, 'existencia': 3.0, 'pedido': 5.0, 'alcanza_meses': 4.0,
             'llegada_dias': 90, 'llegada': '3 meses', 'pedir': 0, 'pedir_ya': False, 'estado': 'alcanza'},
            'alcanza 8 ÷ 2 = 4 meses; tarda 3: pedir = 2 × 3 − 3 − 5 = −2 → 0')
        self.assertEqual((sofa['alcanza'], sofa['proveedor'], sofa['referencia'], sofa['producto']),
                         ('4 meses', self.prov_a.name, 'REP-SOFA', 'Sofá reponer'))
        # Con 2 meses de reserva: 2 × (3 + 2) − 3 − 5 = 2; alcanza (4) cubre la llegada (3): «Pedir», no «Pedir ya».
        sofa = self.fila(self.calcular(reserva=2), self.sofa)
        self.assertEqual((sofa['pedir'], sofa['pedir_ya'], sofa['estado'], sofa['aviso']), (2, False, 'pedir', 'Pedir'))

    def test_pedir_ya_cuando_lo_que_hay_no_cubre_la_llegada(self):
        self.precios_de(self.prov_a).filtered(lambda p: p.product_tmpl_id == self.sofa.product_tmpl_id).delay = 120
        self.venta_facturada(date(2024, 3, 1), self.sofa, 15)                      # 3 al mes
        self.poner_existencia(self.sofa, 2)
        sofa = self.fila(self.calcular(), self.sofa)
        self.assertEqual((sofa['alcanza'], sofa['llegada'], sofa['pedir'], sofa['pedir_ya'], sofa['aviso']),
                         ('20 días', '4 meses', 10, True, 'Pedir ya'), '3 × 4 − 2 = 10; 2 ÷ 3 = 20 días < 4 meses')
        self.assertEqual(self.fila(self.calcular(reserva=1), self.sofa)['pedir'], 13, '3 × (4 + 1) − 2')
        self.assertEqual(self.fila(self.calcular(reserva=0.5), self.sofa)['pedir'], 12, '3 × 4.5 − 2 = 11.5 → 12')
        # Sin nada en la tienda ni en camino: «Ya no hay».
        self.poner_existencia(self.sofa, -2)
        self.assertEqual(self.fila(self.calcular(), self.sofa)['alcanza'], 'Ya no hay')

    def test_sin_tiempo_de_llegada_no_se_calcula_cuanto_pedir(self):
        self.venta_facturada(date(2024, 4, 1), self.mesa, 5)
        self.venta_facturada(date(2024, 4, 1), self.silla, 2)
        filas = self.calcular()
        mesa, silla = self.fila(filas, self.mesa), self.fila(filas, self.silla)
        self.assertEqual(self.mesa.seller_ids.delay, 1, 'Odoo pone 1 día solo')
        self.assertEqual((mesa['llegada_dias'], mesa['llegada'], mesa['pedir'], mesa['pedir_ya'], mesa['aviso']),
                         (None, 'Falta', None, False, 'Falta el tiempo de llegada'))
        self.assertEqual((silla['proveedor_id'], silla['llegada'], silla['pedir'], silla['estado']),
                         (False, 'Sin proveedor', None, 'sin_proveedor'))
        self.assertEqual(mesa['ritmo'], 1.0, 'el ritmo sí se ve')
        # La carga del sistema anterior deja 0 si no venía el plazo: también falta.
        self.mesa.seller_ids.delay = 0
        self.assertIsNone(self.fila(self.calcular(), self.mesa)['pedir'])
        # Un plazo de verdad (más de 1 día) cuenta aunque lo hayan puesto en la ficha del producto.
        self.mesa.seller_ids.delay = 60
        self.assertEqual(self.fila(self.calcular(), self.mesa)['pedir'], 2, '1 × 2 − 0')
        # 1 día puesto a propósito desde Reponer sí cuenta.
        self.env['dcasa.reponer'].with_user(self.gerencia).poner_llegada(self.prov_b.id, 1, 'dias')
        mesa = self.fila(self.calcular(), self.mesa)
        self.assertEqual((mesa['llegada_dias'], mesa['llegada'], mesa['pedir']), (1, '1 día', 1))

    def test_lo_vendido_que_espera_entrega_no_esta_disponible(self):
        """Decisión de la dueña (2026-10-22): lo vendido que sigue en bodega «no está disponible a no ser que se
        revierta la entrega». Disponible = en bodega − lo que falta entregar a clientes; al entregar no se cuenta dos
        veces y al devolver vuelve solo."""
        self.venta_facturada(date(2024, 3, 1), self.sofa, 10)                     # 2 al mes, tarda 3
        self.poner_existencia(self.sofa, 5)
        pedido = self.env['sale.order'].create({'partner_id': self.cliente_rep.id, 'order_line': [
            Command.create({'product_id': self.sofa.id, 'product_uom_qty': 3})]})
        pedido.action_confirm()
        sofa = self.fila(self.calcular(), self.sofa)
        self.assertEqual((sofa['en_bodega'], sofa['por_entregar'], sofa['existencia'], sofa['alcanza'], sofa['pedir']),
                         (5.0, 3.0, 2.0, '1 mes', 4), 'disponible 5 − 3 = 2; pedir 2 × 3 − 2')
        # Un pedido sin confirmar no compromete nada; uno cancelado tampoco.
        borrador = self.env['sale.order'].create({'partner_id': self.cliente_rep.id, 'order_line': [
            Command.create({'product_id': self.sofa.id, 'product_uom_qty': 9})]})
        self.assertEqual(self.fila(self.calcular(), self.sofa)['por_entregar'], 3.0)
        borrador.action_confirm()
        borrador._action_cancel()
        self.assertEqual(self.fila(self.calcular(), self.sofa)['por_entregar'], 3.0)
        # Entregado: sale de la bodega y deja de estar por entregar (no se resta dos veces).
        entrega = pedido.picking_ids
        entrega.move_ids.quantity = 3
        entrega.button_validate()
        sofa = self.fila(self.calcular(), self.sofa)
        self.assertEqual((sofa['en_bodega'], sofa['por_entregar'], sofa['existencia']), (2.0, 0.0, 2.0))
        # Devuelto: vuelve a la bodega y cuenta solo.
        devolucion = self.env['stock.return.picking'].with_context(
            active_id=entrega.id, active_model='stock.picking').create({})
        devolucion.product_return_moves.quantity = 3
        regreso = devolucion._create_return()
        regreso.move_ids.quantity = 3
        regreso.button_validate()
        self.assertEqual(self.fila(self.calcular(), self.sofa)['existencia'], 5.0, 'la devolución vuelve a contar')
        # Vendido más de lo que hay: sale en negativo y suma a lo que hay que pedir.
        mas = self.env['sale.order'].create({'partner_id': self.cliente_rep.id, 'order_line': [
            Command.create({'product_id': self.sofa.id, 'product_uom_qty': 7})]})
        mas.action_confirm()
        sofa = self.fila(self.calcular(), self.sofa)
        self.assertEqual((sofa['existencia'], sofa['alcanza'], sofa['pedir'], sofa['pedir_ya']),
                         (-2.0, 'Ya no hay', 8, True), '2 × 3 − (−2) = 8')

    def test_la_ventana_de_meses_que_se_elige(self):
        self.venta_facturada(date(2023, 12, 31), self.lampara, 6)                 # fuera de 5 meses, dentro de 6
        self.venta_facturada(date(2024, 1, 1), self.lampara, 4)                   # primer día de la ventana de 5
        self.venta_facturada(date(2024, 5, 31), self.lampara, 1)                  # último día
        self.venta_facturada(date(2024, 6, 1), self.lampara, 50)                  # después: no cuenta
        self.assertEqual(self.env['dcasa.reponer']._ventana(5, HASTA), (date(2024, 1, 1), HASTA))
        cinco = self.fila(self.calcular(5), self.lampara)
        self.assertEqual((cinco['vendido'], cinco['ritmo']), (5.0, 1.0))
        seis = self.fila(self.calcular(6), self.lampara)
        self.assertEqual((seis['vendido'], seis['ritmo']), (11.0, 1.83))
        self.assertEqual(self.fila(self.calcular(1), self.lampara)['vendido'], 1.0, 'solo mayo')
        self.assertEqual(self.fila(self.calcular(12), self.lampara)['vendido'], 11.0)
        for meses, reserva in ((0, 0), (13, 0), ('x', 0), (5, -1), (5, 13), (5, 'nan')):
            with self.subTest(meses=meses, reserva=reserva), self.assertRaises(UserError):
                self.calcular(meses, reserva)

    def test_las_notas_de_credito_restan_y_el_flete_no_cuenta(self):
        factura = self.suelta(date(2024, 5, 10), lineas=[
            Command.create({'product_id': self.sofa.id, 'quantity': 3, 'price_unit': 400.0}),
            Command.create({'product_id': self.flete.id, 'quantity': 1, 'price_unit': 25.0})])
        nota = self.nota(factura, date(2024, 5, 12), precio=400.0)          # devuelve 1 sofá (su primera línea)
        self.assertEqual(nota.invoice_line_ids.product_id, self.sofa)
        filas = self.calcular()
        self.assertEqual(self.fila(filas, self.sofa)['vendido'], 2.0)
        self.assertNotIn(self.flete.id, [fila['producto_id'] for fila in filas])
        tablero = self.env['dcasa.tablero']
        unidades = tablero._unidades_por_producto(tablero._filas(date(2024, 1, 1), HASTA, 'facturado'))
        self.assertEqual(unidades[self.sofa][0], 2.0, 'las mismas unidades de «Lo más vendido» del Inicio')
        # Vendido y devuelto entero: no sale.
        self.nota(self.venta_facturada(date(2024, 5, 13), self.mesa, 1), date(2024, 5, 14))
        self.assertNotIn(self.mesa.id, [fila['producto_id'] for fila in self.calcular()])

    def test_orden_por_urgencia_y_grupos_por_proveedor_con_totales(self):
        self.venta_facturada(date(2024, 5, 1), self.sofa, 10)      # 2/mes, hay 0, tarda 3 → pedir ya, pedir 6
        self.venta_facturada(date(2024, 5, 1), self.lampara, 5)    # 1/mes, hay 10 → alcanza
        self.poner_existencia(self.lampara, 10)
        self.venta_facturada(date(2024, 5, 1), self.mesa, 5)       # falta el tiempo de llegada
        self.venta_facturada(date(2024, 5, 1), self.silla, 5)      # sin proveedor
        filas = self.calcular()
        self.assertEqual([fila['producto_id'] for fila in filas],
                         [self.sofa.id, self.mesa.id, self.silla.id, self.lampara.id])
        self.assertEqual([fila['estado'] for fila in filas], ['ya', 'falta', 'sin_proveedor', 'alcanza'])
        grupos = self.env['dcasa.reponer'].with_user(self.gerencia)._grupos(filas)
        self.assertEqual([grupo['proveedor_id'] for grupo in grupos], [self.prov_a.id, self.prov_b.id, False])
        a = grupos[0]
        self.assertEqual((a['ya'], a['vendido'], a['ritmo'], a['existencia'], a['pedido'], a['pedir']),
                         (1, 15.0, 3.0, 10.0, 0.0, 6))
        self.assertEqual(grupos[2]['proveedor'], 'Sin proveedor')
        self.assertEqual([grupo['calculables'] for grupo in grupos], [2, 0, 0], 'sin llegada, el total de pedir es «—»')

    def test_filtro_por_proveedor_y_lo_que_no_se_vendio(self):
        self.venta_facturada(date(2024, 5, 10), self.sofa, 2)
        self.venta_facturada(date(2024, 5, 10), self.mesa, 1)
        self.assertEqual([fila['producto_id'] for fila in self.calcular(proveedor=self.prov_a)], [self.sofa.id])
        self.assertEqual([fila['producto_id'] for fila in self.calcular(proveedor=self.prov_b)], [self.mesa.id],
                         'el sofá también tiene precio de B, pero su principal es A')
        Reponer = self.env['dcasa.reponer'].with_user(self.gerencia)
        self.assertEqual(Reponer._sin_ventas(self.prov_a, self.calcular(proveedor=self.prov_a)), 1, 'la lámpara')
        self.assertEqual(Reponer._sin_ventas(self.prov_b, self.calcular(proveedor=self.prov_b)), 0,
                         'del sofá, B no es el principal')

    # --- La pantalla, el buscador y el tiempo de llegada ------------------------------------------------------------

    def test_la_pantalla_trae_ventana_filas_grupos_y_resumen(self):
        hoy = self.env['dcasa.reponer'].with_user(self.gerencia)._ventana(5)[1]
        self.venta_facturada(hoy - timedelta(days=3), self.sofa, 10)
        self.venta_facturada(hoy - timedelta(days=3), self.mesa, 2)
        datos = self.env['dcasa.reponer'].with_user(self.gerencia).obtener(5, 0, self.prov_a.id)
        self.assertEqual((datos['meses'], datos['reserva'], datos['hasta']), (5, 0.0, hoy.strftime('%d/%m/%Y')))
        self.assertEqual([fila['producto_id'] for fila in datos['filas']], [self.sofa.id])
        self.assertEqual(datos['proveedor']['nombre'], self.prov_a.name)
        self.assertEqual(datos['sin_ventas'], 1, 'la lámpara no se vendió')
        self.assertEqual(datos['grupos'][0]['producto_ids'], [self.sofa.id])
        self.assertEqual(datos['grupos'][0]['llegada']['texto'], '3 meses')
        self.assertEqual((datos['resumen']['ya'], datos['unidades']), (1, 6))
        todos = self.env['dcasa.reponer'].with_user(self.gerencia).obtener()
        self.assertIsNone(todos['proveedor'])
        self.assertEqual(todos['meses'], 5, 'de fábrica, 5 meses')
        self.assertLessEqual({self.sofa.id, self.mesa.id}, {fila['producto_id'] for fila in todos['filas']})
        bambu = self.env['dcasa.reponer'].with_user(self.gerencia).obtener(proveedor_id=self.prov_b.id)
        self.assertEqual((bambu['resumen']['falta'], bambu['unidades'], bambu['grupos'][0]['calculables']), (1, 0, 0),
                         'la mesa: falta el tiempo de llegada')

    def test_buscador_de_proveedores(self):
        Socio = self.env['res.partner']
        cliente = Socio.create({'name': 'CLIENTE DE PRUEBA que no es proveedor', 'phone': '+507 6111-2239'})
        self.compra_al_proveedor(self.prov_a, self.sofa, 1)
        Reponer = self.env['dcasa.reponer'].with_user(self.gerencia)

        def fichas(texto):
            return {ficha['id']: ficha for ficha in Reponer.buscar_proveedores(texto)['proveedores']}

        for texto in ('asia', 'PRUEBA Ás', '6111-2233', '61112233', '155000111', 'REP-LAMP', 'lámpara reponer'):
            with self.subTest(texto=texto):
                encontrados = fichas(texto)
                self.assertIn(self.prov_a.id, encontrados)
                self.assertNotIn(self.prov_b.id, encontrados)
        self.assertIn(self.prov_a.id, fichas('6111-223'))
        self.assertNotIn(cliente.id, fichas('6111-223'), 'el teléfono también es de un cliente, que no es proveedor')
        self.assertIn(self.prov_b.id, fichas('bambu'))
        self.assertLessEqual({self.prov_a.id, self.prov_b.id}, set(fichas('REP-SOFA')), 'los dos venden el sofá')
        self.assertFalse(fichas('ningún proveedor se llama así'))
        ficha = fichas('asia')[self.prov_a.id]
        self.assertEqual((ficha['productos'], ficha['telefono'], ficha['ruc'], ficha['coincide']),
                         (2, '+507 6111-2233', '155000111-2-2020', ''))
        self.assertRegex(ficha['ultima_compra'], r'^\d{2}/\d{2}/\d{4}$')
        self.assertEqual(ficha['llegada'], {'dias': 90, 'texto': '3 meses', 'falta': 0, 'precios': 2})
        self.assertEqual(fichas('REP-LAMP')[self.prov_a.id]['coincide'], 'Vende «Lámpara reponer»')
        self.assertEqual(fichas('6111 2233')[self.prov_a.id]['coincide'], 'Teléfono +507 6111-2233')
        self.assertEqual(fichas('155000111')[self.prov_a.id]['coincide'], 'RUC 155000111-2-2020')
        b = fichas('')[self.prov_b.id]
        self.assertEqual((b['ultima_compra'], b['llegada']['texto'], b['llegada']['falta']),
                         ('', 'Falta el tiempo de llegada', 2))

    def test_gerencia_pone_el_tiempo_de_llegada(self):
        Reponer = self.env['dcasa.reponer'].with_user(self.gerencia)
        ficha = Reponer.poner_llegada(self.prov_b.id, '3', 'meses')
        precios = self.precios_de(self.prov_b)
        self.assertEqual(len(precios), 2, 'el de la mesa y el segundo del sofá')
        self.assertEqual(set(precios.mapped('delay')), {90})
        self.assertTrue(all(precios.mapped('dcasa_llegada_puesta')))
        self.assertEqual(ficha['llegada'], {'dias': 90, 'texto': '3 meses', 'falta': 0, 'precios': 2})
        self.assertIn('3 meses (90 días)', str(self.prov_b.message_ids[0].body))
        self.assertEqual(Reponer.poner_llegada(self.prov_b.id, 45, 'dias')['llegada']['texto'], '45 días')
        self.assertEqual(Reponer.poner_llegada(self.prov_b.id, '3,5', 'meses')['llegada']['dias'], 105)
        for cantidad, unidad in ((0, 'meses'), (-1, 'dias'), ('abc', 'dias'), (25, 'meses'), (731, 'dias'),
                                 (3, 'semanas'), ('inf', 'dias')):
            with self.subTest(cantidad=cantidad, unidad=unidad), self.assertRaises(UserError):
                Reponer.poner_llegada(self.prov_b.id, cantidad, unidad)
        self.assertEqual(set(precios.mapped('delay')), {105}, 'lo inválido no cambia nada')
        sin_precios = self.env['res.partner'].create({'name': 'PROVEEDOR DE PRUEBA sin productos'})
        with self.assertRaisesRegex(UserError, 'todavía no tiene productos'):
            Reponer.poner_llegada(sin_precios.id, 3, 'meses')

    def test_la_vendedora_no_puede_nada_de_reponer(self):
        menu = self.env.ref('dcasa_interfaz.menu_dcasa_reponer')
        self.assertIn(menu.id, self.env['ir.ui.menu'].with_user(self.gerencia)._visible_menu_ids())
        self.assertNotIn(menu.id, self.env['ir.ui.menu'].with_user(self.vendedora)._visible_menu_ids())
        self.assertEqual(self.env.ref('dcasa_interfaz.action_dcasa_reponer').tag, 'dcasa_reponer')
        Reponer = self.env['dcasa.reponer'].with_user(self.vendedora)
        for que, llamada in (('ver', Reponer.obtener), ('buscar', lambda: Reponer.buscar_proveedores('asia')),
                             ('poner', lambda: Reponer.poner_llegada(self.prov_b.id, 3, 'meses')),
                             ('excel', Reponer._excel), ('calcular', lambda: self.calcular(usuario=self.vendedora))):
            with self.subTest(que), self.assertRaises(AccessError):
                llamada()
        self.assertEqual(set(self.precios_de(self.prov_b).mapped('delay')), {1}, 'nada cambió')
        self.assertFalse(any(self.precios_de(self.prov_b).mapped('dcasa_llegada_puesta')))

    def test_el_excel_trae_lo_mismo_que_la_pantalla(self):
        hoy = self.env['dcasa.reponer'].with_user(self.gerencia)._ventana(5)[1]
        self.venta_facturada(hoy - timedelta(days=5), self.sofa, 10)       # 2/mes, hay 0 → pedir ya
        self.venta_facturada(hoy - timedelta(days=5), self.lampara, 5)     # 1/mes, hay 10 → alcanza
        self.poner_existencia(self.lampara, 10)
        Reponer = self.env['dcasa.reponer'].with_user(self.gerencia)
        nombre, contenido = Reponer._excel(5, 1, self.prov_a.id)
        self.assertEqual(nombre, f'reponer_proveedor-de-prueba-asia_{hoy.isoformat()}.xlsx')
        libro = openpyxl.load_workbook(io.BytesIO(contenido))
        self.assertEqual(libro.sheetnames, ['Reponer', 'Por proveedor'])
        hoja = libro['Reponer']
        self.assertEqual(hoja['A1'].value, 'Reponer · PROVEEDOR DE PRUEBA Ásia')
        self.assertIn('Reserva: 1 meses', hoja['A2'].value)
        self.assertEqual([celda.value for celda in hoja[5]], [
            'Aviso', 'Producto', 'Referencia', 'Proveedor', 'Vendido en 5 meses', 'Se vende al mes', 'En bodega',
            'Vendido por entregar', 'Disponible hoy (sin lo ya vendido)', 'Viene en camino', 'Alcanza para (meses)',
            'Tarda en llegar (días)', 'Pedir'])
        filas = [[celda.value for celda in fila] for fila in hoja.iter_rows(min_row=6)]
        self.assertEqual([fila[1] for fila in filas],
                         [fila['producto'] for fila in Reponer._calcular(5, 1, self.prov_a)], 'el mismo orden')
        self.assertEqual(filas[0], ['Pedir ya', 'Sofá reponer', 'REP-SOFA', self.prov_a.name, 10, 2, 0, 0, 0, 0, 0,
                                    90, 8], '2 × (3 + 1) − 0 − 0 = 8')
        self.assertEqual(filas[1], ['Alcanza', 'Lámpara reponer', 'REP-LAMP', self.prov_a.name, 5, 1, 10, 0, 10, 0, 10,
                                    90, 0], '1 × (3 + 1) − 10 < 0')
        resumen = libro['Por proveedor']
        self.assertEqual([celda.value for celda in resumen[2]], [self.prov_a.name, '3 meses', 2, 1, 15, 3, 10, 0, 8])
        self.assertEqual((resumen['A3'].value, resumen['I3'].value), ('Total', '=SUM(I2:I2)'))


@tagged('post_install', '-at_install')
class TestReponerEnElNavegador(ReponerDatos, VentasMixin, HttpCase):
    """La pantalla de verdad en Chrome, a 1440 y a 390 px: buscar el proveedor (por código y por producto), elegirlo,
    ver «Pedir ya» primero, la reserva, agrupar por proveedor, poner el tiempo de llegada que faltaba y el Excel (que
    la vendedora no baja)."""

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.crear_datos()
        hoy = cls.env['dcasa.reponer'].with_user(cls.gerencia)._ventana(5)[1]
        cls.venta_facturada(hoy - timedelta(days=4), cls.sofa, 10)       # A: 2/mes, hay 0 → pedir ya, pedir 6
        cls.venta_facturada(hoy - timedelta(days=4), cls.lampara, 5)     # A: alcanza
        cls.poner_existencia(cls.lampara, 30)
        cls.venta_facturada(hoy - timedelta(days=4), cls.mesa, 5)        # B: falta el tiempo de llegada
        cls.mesa.seller_ids.product_code = 'BAMBU-77'                    # su código en el proveedor

    def test_el_excel_solo_para_gerencia(self):
        self.authenticate('gerencia_rol', 'gerencia_rol')
        respuesta = self.url_open(f'/dcasa/reponer/excel?meses=5&reserva=0&proveedor={self.prov_a.id}')
        self.assertEqual(respuesta.status_code, 200)
        self.assertEqual(respuesta.headers['Content-Type'],
                         'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        self.assertIn('attachment', respuesta.headers['Content-Disposition'])
        productos = [celda.value for celda in openpyxl.load_workbook(io.BytesIO(respuesta.content))['Reponer']['B']]
        self.assertIn('Sofá reponer', productos)
        self.authenticate('vendedora_rol', 'vendedora_rol')
        with mute_logger('odoo.http'):
            self.assertEqual(self.url_open('/dcasa/reponer/excel?meses=5').status_code, 403)

    def test_la_pantalla_a_1440_y_a_390(self):
        codigo = '(async () => {' + ESPERAR_JS + r"""
            const raiz = () => document.querySelector('.o_dcasa_reponer');
            const sinScrollDeCostado = (que) => {
                const ancho = document.documentElement.clientWidth;
                for (const el of [raiz(), ...raiz().querySelectorAll('.o_dcasa_reponer_paso, tr, button, a.btn')]) {
                    const caja = el.getBoundingClientRect();
                    if (caja.width && (caja.right > ancho + 1 || caja.left < -1)) {
                        throw new Error(`se sale (${que}): ${el.className} ${caja.left}–${caja.right} > ${ancho}`);
                    }
                }
                if (raiz().scrollWidth > raiz().clientWidth + 1) {
                    throw new Error(`scroll de costado (${que}): ${raiz().scrollWidth} > ${raiz().clientWidth}`);
                }
            };
            const escribir = (input, texto) => {
                input.value = texto;
                input.dispatchEvent(new Event('input', { bubbles: true }));
            };
            const cambiar = (select, valor) => {
                select.value = valor;
                select.dispatchEvent(new Event('change', { bubbles: true }));
            };
            const pedirDe = (fila) => fila.querySelector('.o_dcasa_reponer_pedir').textContent.trim();
            await esperar(() => raiz() && document.querySelector('.o_dcasa_reponer_todos'), 'la pantalla');
            // La explicación: abierta en la computadora; en el celular, cerrada (no tapa la pantalla) y se abre.
            const como = document.querySelector('.o_dcasa_reponer_como');
            const celular = document.documentElement.clientWidth < 768;
            igual(como.open, !celular, 'la explicación abierta solo en la computadora');
            if (celular) {
                como.querySelector('summary').click();
                await esperar(() => como.open, 'se abre al tocarla');
            }
            if (!document.querySelector('.o_dcasa_reponer_ejemplo').textContent.includes('números inventados')) {
                throw new Error('el ejemplo dice que es ejemplo');
            }
            igual(document.querySelector('.o_dcasa_reponer_meses').value, '5', 'de fábrica, 5 meses');
            igual(document.querySelector('.o_dcasa_reponer_reserva').value, '0', 'de fábrica, sin reserva');
            sinScrollDeCostado('al abrir');
            // Paso 1: el buscador, con su instrucción. Por el código del proveedor y por parte de un producto.
            const buscar = document.querySelector('#o_dcasa_reponer_buscar');
            const instruccion = document.querySelector('label[for="o_dcasa_reponer_buscar"]').textContent;
            for (const frase of ['nombre', 'teléfono', 'RUC', 'producto']) {
                if (!instruccion.includes(frase)) { throw new Error('el buscador explica: ' + frase); }
            }
            escribir(buscar, 'BAMBU-77');
            await esperar(() => textos('.o_dcasa_reponer_opcion strong').includes('PROVEEDOR DE PRUEBA Bambú')
                && !textos('.o_dcasa_reponer_opcion strong').includes('PROVEEDOR DE PRUEBA Ásia'), 'por código');
            escribir(buscar, 'lámpara repo');
            const opcion = await esperar(() => [...document.querySelectorAll('.o_dcasa_reponer_opcion')]
                .find((b) => b.textContent.includes('PROVEEDOR DE PRUEBA Ásia')), 'el proveedor del producto');
            for (const frase of ['2 productos', 'Tarda 3 meses', 'Vende «Lámpara reponer»']) {
                if (!opcion.textContent.includes(frase)) { throw new Error('la opción debe decir: ' + frase); }
            }
            if (opcion.getBoundingClientRect().height < 44) { throw new Error('la opción es chica para el dedo'); }
            sinScrollDeCostado('con resultados');
            opcion.click();
            await esperar(() => document.querySelector('.o_dcasa_reponer_elegido')
                && document.querySelectorAll('.o_dcasa_reponer_fila').length === 2, 'la lista del proveedor');
            igual(document.querySelector('.o_dcasa_reponer_elegido strong').textContent, 'PROVEEDOR DE PRUEBA Ásia',
                  'el elegido');
            // «Pedir ya» primero (el sofá: 2 × 3 − 0 = 6), después la lámpara, que alcanza.
            const filas = [...document.querySelectorAll('.o_dcasa_reponer_fila')];
            igual(filas.map((f) => f.querySelector('.o_dcasa_reponer_aviso').textContent.trim()),
                  ['Pedir ya', 'Alcanza'], 'el orden por urgencia');
            igual(pedirDe(filas[0]), '6', 'pedir 2 × 3 − 0');
            const ya = filas[0].querySelector('.o_aviso_ya');
            igual([getComputedStyle(ya).color, getComputedStyle(ya).backgroundColor],
                  ['rgb(254, 208, 15)', 'rgb(11, 31, 82)'], 'el amarillo, solo sobre azul oscuro');
            if (document.querySelector('.o_dcasa_reponer_sin_ventas')) {
                throw new Error('los dos productos del proveedor se vendieron');
            }
            // El Excel: un enlace que baja el archivo hecho en el servidor.
            const excel = document.querySelector('a.o_dcasa_reponer_excel');
            const url = new URL(excel.href);
            igual([url.pathname, url.searchParams.get('meses'), url.searchParams.get('reserva'),
                   url.searchParams.has('proveedor')], ['/dcasa/reponer/excel', '5', '0', true], 'el enlace del Excel');
            const bajado = await fetch(excel.href);
            igual([bajado.status, bajado.headers.get('content-type')],
                  [200, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'], 'el Excel baja');
            sinScrollDeCostado('con la lista');
            // Paso 2: 1 mes de reserva → 2 × (3 + 1) − 0 = 8.
            cambiar(document.querySelector('.o_dcasa_reponer_reserva'), '1');
            await esperar(() => pedirDe(document.querySelector('.o_dcasa_reponer_fila')) === '8', 'la reserva');
            cambiar(document.querySelector('.o_dcasa_reponer_reserva'), '0');
            await esperar(() => pedirDe(document.querySelector('.o_dcasa_reponer_fila')) === '6', 'sin reserva');
            // Todos los proveedores, por proveedor: la mesa de Bambú no tiene tiempo de llegada.
            document.querySelector('.o_dcasa_reponer_cambiar').click();
            (await esperar(() => document.querySelector('.o_dcasa_reponer_todos'), 'el buscador otra vez')).click();
            await esperar(() => textos('.o_dcasa_reponer_nombre').includes('Mesa reponer'), 'todos');
            document.querySelector('.o_vista_proveedor').click();
            const cabezas = await esperar(() => {
                const t = textos('.o_dcasa_reponer_grupo_cabeza strong')
                    .filter((nombre) => nombre.startsWith('PROVEEDOR DE PRUEBA'));
                return t.length === 2 && t;
            }, 'los grupos');
            igual(cabezas, ['PROVEEDOR DE PRUEBA Ásia', 'PROVEEDOR DE PRUEBA Bambú'], 'primero el del «Pedir ya»');
            const grupoDe = (nombre) => [...document.querySelectorAll('.o_dcasa_reponer_grupo')]
                .find((g) => g.querySelector('.o_dcasa_reponer_grupo_cabeza strong').textContent === nombre);
            const filaDe = (grupo, nombre) => [...grupo.querySelectorAll('.o_dcasa_reponer_fila')]
                .find((f) => f.textContent.includes(nombre));
            const grupoA = grupoDe('PROVEEDOR DE PRUEBA Ásia');
            igual(pedirDe(grupoA.querySelector('.o_dcasa_reponer_total')), '6', 'el total de pedir del grupo');
            const grupoB = grupoDe('PROVEEDOR DE PRUEBA Bambú');
            const mesa = filaDe(grupoB, 'Mesa reponer');
            igual(pedirDe(mesa), '—', 'sin tiempo de llegada no hay cuánto pedir');
            igual(pedirDe(grupoB.querySelector('.o_dcasa_reponer_total')), '—', 'ni total de cuánto pedir');
            if (!mesa.textContent.includes('Falta el tiempo de llegada')) { throw new Error('la fila dice que falta'); }
            sinScrollDeCostado('por proveedor');
            // Poner el tiempo de llegada de Bambú: 2 meses.
            grupoB.querySelector('.o_dcasa_reponer_poner').click();
            const ventana = await esperar(() => document.querySelector('.modal .o_dcasa_reponer_llegada_form'),
                                          'la ventana del tiempo de llegada');
            const cantidad = ventana.querySelector('.o_dcasa_reponer_cantidad');
            cantidad.value = '2';
            cantidad.dispatchEvent(new Event('input', { bubbles: true }));
            await esperar(() => (document.querySelector('.o_dcasa_reponer_equivale') || {}).textContent === '= 60 días',
                          'dice cuántos días son');
            document.querySelector('.modal .o_dcasa_reponer_guardar').click();
            await esperar(() => !document.querySelector('.modal .o_dcasa_reponer_llegada_form'), 'se cierra');
            await esperar(() => {
                const fila = filaDe(grupoDe('PROVEEDOR DE PRUEBA Bambú'), 'Mesa reponer');
                return fila && pedirDe(fila) === '2';
            }, 'la mesa ya calcula: 1 × 2 − 0 = 2');
            sinScrollDeCostado('después de poner el tiempo');
        """ + "console.log('test successful'); })();"
        for tamano in ('1440x900', '390x844'):
            with self.subTest(tamano=tamano):
                self.precios_de(self.prov_b).write({'delay': 1, 'dcasa_llegada_puesta': False})
                self.browser_size = tamano
                self.browser_js('/odoo/action-dcasa_interfaz.action_dcasa_reponer', codigo, login='gerencia_rol',
                                timeout=180)
                self.assertEqual(set(self.precios_de(self.prov_b).mapped('delay')), {60})
