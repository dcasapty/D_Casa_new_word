"""Prototipo (ronda 2, brian-excel): extractor determinista de .xlsx -> esquema intermedio JSON.

No usa ningún modelo de IA ni toca Odoo. Solo biblioteca estándar + openpyxl (+ Pillow opcional).
Uso:
    python extractor.py ARCHIVO.xlsx --salida salida/ [--fotos "up media"]

Produce en --salida:
    <nombre>.intermedio.json   celda -> valor/fórmula/formato/imagen/comentario, por hoja
    <nombre>.resumen.json      números medidos (hojas, combinadas, fórmulas, imágenes, riesgos)
    <nombre>.candidatos.json   filas de la tabla principal normalizadas + discrepancias
                               (determinista: precios «Full $149.99 · Queen $169.99» con regex)

Qué lee (y por qué importa):
  * valores Y fórmulas (se abre el libro dos veces: data_only=False/True). Un valor sin
    caché (archivo generado por script, nunca abierto en Excel) sale como `valor: null`
    con `formula` presente -> se marca en `riesgos`.
  * celdas combinadas: el valor vive en la celda superior izquierda; las demás se marcan
    `combinada_en` para no leerlas como vacías.
  * imágenes FLOTANTES (xl/drawings): ancla `twoCell/oneCell/absolute` -> celda de origen,
    hash sha256 del binario (para deduplicar) y tamaño. Se leen del XML crudo, no de
    openpyxl, porque openpyxl pierde formas y no expone el ancla final.
  * imágenes EN CELDA («Colocar en celda» / =IMAGEN()): cadena richData
    (celda vm -> metadata.xml -> rdrichvalue.xml -> richValueRel.xml -> media). openpyxl
    NO las ve; aquí se resuelven a mano (formato no documentado en ECMA-376: NO VERIFICADO
    contra todas las versiones de Excel, probado solo con el archivo sintético de pruebas/).
  * comentarios/notas, hipervínculos, hojas y filas/columnas ocultas, rangos con nombre,
    filtros, validaciones de datos.
  * el contenido de celdas se trata como DATO: se guarda tal cual y se marca
    `sospecha_instruccion` si parece una orden para el modelo (defensa B-07/S-10).
"""
from __future__ import annotations

import argparse
import hashlib
import json
import posixpath
import re
import sys
import zipfile
from collections import Counter, OrderedDict
from pathlib import Path
from xml.etree import ElementTree as ET

import openpyxl
from openpyxl.utils import get_column_letter

ESQUEMA = 'dcasa.xlsx/1'
NS = {
    'm': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main',
    'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
    'xdr': 'http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing',
    'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
    'rel': 'http://schemas.openxmlformats.org/package/2006/relationships',
    'rv': 'http://schemas.microsoft.com/office/spreadsheetml/2017/richdata',
    'rvrel': 'http://schemas.microsoft.com/office/spreadsheetml/2022/richvaluerel',
}
LIMITE_ZIP_MB = 200          # tamaño descomprimido máximo (defensa contra zip bomb)
LIMITE_RATIO = 100           # descomprimido / comprimido
PATRON_ORDEN = re.compile(
    r'(ignora|olvida|ignore|disregard|system\s*:|eres ahora|you are now|llama a la herramienta|'
    r'cambia el rol|elimina|borra todo)', re.I)
TAMANOS = ('Twin', 'Full', 'Queen', 'King')


# ----------------------------------------------------------------------------------------
# Seguridad del contenedor
# ----------------------------------------------------------------------------------------
def revisar_zip(ruta: Path) -> list[str]:
    """Rechaza zip bombs y rutas peligrosas ANTES de abrir con openpyxl."""
    avisos = []
    with zipfile.ZipFile(ruta) as z:
        total = sum(i.file_size for i in z.infolist())
        comprimido = max(1, sum(i.compress_size for i in z.infolist()))
        if total > LIMITE_ZIP_MB * 1024 * 1024:
            raise ValueError(f'Descomprimido {total/1e6:.0f} MB > {LIMITE_ZIP_MB} MB: se rechaza.')
        if total / comprimido > LIMITE_RATIO:
            raise ValueError(f'Razón de compresión {total/comprimido:.0f}x sospechosa: se rechaza.')
        for i in z.infolist():
            if i.filename.startswith('/') or '..' in i.filename.split('/'):
                raise ValueError(f'Ruta peligrosa en el zip: {i.filename}')
        nombres = z.namelist()
        if 'xl/vbaProject.bin' in nombres:
            avisos.append('macros_vba_presentes (no se ejecutan; se ignoran)')
        if any(n.startswith('xl/externalLinks/') for n in nombres):
            avisos.append('vinculos_externos_presentes (no se siguen)')
    return avisos


# ----------------------------------------------------------------------------------------
# Imágenes: flotantes (drawings) y en celda (richData)
# ----------------------------------------------------------------------------------------
def _rels(z: zipfile.ZipFile, ruta_parte: str) -> dict[str, tuple[str, str]]:
    """Relaciones de una parte: {rId: (tipo, destino_absoluto_en_zip)}."""
    carpeta, nombre = posixpath.split(ruta_parte)
    ruta_rels = posixpath.join(carpeta, '_rels', nombre + '.rels')
    if ruta_rels not in z.namelist():
        return {}
    raiz = ET.fromstring(z.read(ruta_rels))
    salida = {}
    for r in raiz.findall('rel:Relationship', NS):
        destino = r.get('Target')
        if r.get('TargetMode') == 'External':
            continue
        absoluto = posixpath.normpath(destino[1:] if destino.startswith('/') else posixpath.join(carpeta, destino))
        salida[r.get('Id')] = (r.get('Type', '').rsplit('/', 1)[-1], absoluto)
    return salida


def _hoja_a_partes(z: zipfile.ZipFile) -> dict[str, str]:
    """nombre de hoja -> ruta de su parte XML."""
    wb = ET.fromstring(z.read('xl/workbook.xml'))
    rels = _rels(z, 'xl/workbook.xml')
    salida = {}
    for h in wb.find('m:sheets', NS):
        rid = h.get(f'{{{NS["r"]}}}id')
        salida[h.get('name')] = rels[rid][1]
    return salida


def _info_imagen(z: zipfile.ZipFile, ruta_media: str) -> dict:
    datos = z.read(ruta_media)
    info = {'archivo_en_xlsx': ruta_media, 'bytes': len(datos),
            'sha256': hashlib.sha256(datos).hexdigest()[:16]}
    try:
        from io import BytesIO
        from PIL import Image
        with Image.open(BytesIO(datos)) as im:
            info.update(formato=im.format, ancho=im.width, alto=im.height)
    except Exception:  # noqa: BLE001 — Pillow ausente o imagen no estándar (emf/wmf)
        info['formato'] = Path(ruta_media).suffix.lstrip('.').lower()
    return info


def _celda(col0: int, fila0: int) -> str:
    return f'{get_column_letter(col0 + 1)}{fila0 + 1}'


def imagenes_flotantes(z: zipfile.ZipFile, parte_hoja: str) -> list[dict]:
    salida = []
    for tipo, destino in _rels(z, parte_hoja).values():
        if tipo != 'drawing':
            continue
        dibujo = ET.fromstring(z.read(destino))
        rel_dib = _rels(z, destino)
        for ancla in list(dibujo):
            etiqueta = ancla.tag.split('}')[1]          # twoCellAnchor | oneCellAnchor | absoluteAnchor
            if etiqueta not in ('twoCellAnchor', 'oneCellAnchor', 'absoluteAnchor'):
                continue
            pic = ancla.find('xdr:pic', NS)
            if pic is None:
                salida.append({'tipo': 'forma_no_imagen', 'ancla': etiqueta})   # gráfico, forma, texto
                continue
            blip = pic.find('.//a:blip', NS)
            rid = blip.get(f'{{{NS["r"]}}}embed') if blip is not None else None
            if rid not in rel_dib:
                continue
            desde = ancla.find('xdr:from', NS)
            hasta = ancla.find('xdr:to', NS)
            entrada = {'tipo': 'flotante', 'ancla': etiqueta,
                       'nombre': (pic.find('xdr:nvPicPr/xdr:cNvPr', NS).get('name')
                                  if pic.find('xdr:nvPicPr/xdr:cNvPr', NS) is not None else None)}
            if desde is not None:
                entrada['celda_desde'] = _celda(int(desde.find('xdr:col', NS).text), int(desde.find('xdr:row', NS).text))
            if hasta is not None:
                entrada['celda_hasta'] = _celda(int(hasta.find('xdr:col', NS).text), int(hasta.find('xdr:row', NS).text))
            entrada.update(_info_imagen(z, rel_dib[rid][1]))
            salida.append(entrada)
    return salida


def imagenes_en_celda(z: zipfile.ZipFile, parte_hoja: str) -> list[dict]:
    """Celdas con vm="N" -> richData -> media. Devuelve [{celda, ...info_imagen}]."""
    nombres = set(z.namelist())
    if 'xl/metadata.xml' not in nombres or 'xl/richData/rdrichvalue.xml' not in nombres:
        return []
    # 1) vm (1-based) -> índice de valor rico, vía metadata.xml (valueMetadata -> bk -> rc v=)
    meta = ET.fromstring(z.read('xl/metadata.xml'))
    ns_m = NS['m']
    futuros = meta.findall(f'{{{ns_m}}}futureMetadata')
    rv_por_bk = []                                   # índice de bk en futureMetadata -> rvb i
    for fm in futuros:
        if fm.get('name') != 'XLRICHVALUE':
            continue
        for bk in fm.findall(f'{{{ns_m}}}bk'):
            rvb = bk.find('.//{*}rvb')
            rv_por_bk.append(int(rvb.get('i')) if rvb is not None else None)
    vm_a_rv = {}
    vmeta = meta.find(f'{{{ns_m}}}valueMetadata')
    if vmeta is not None:
        for n, bk in enumerate(vmeta.findall(f'{{{ns_m}}}bk'), start=1):
            rc = bk.find(f'{{{ns_m}}}rc')
            if rc is not None and rc.get('v') is not None and int(rc.get('v')) < len(rv_por_bk):
                vm_a_rv[n] = rv_por_bk[int(rc.get('v'))]
    # 2) valor rico -> índice de relación (_localImageIdentifier)
    rv = ET.fromstring(z.read('xl/richData/rdrichvalue.xml'))
    rv_a_rel = {}
    for n, nodo in enumerate(rv.findall('{*}rv')):
        valores = [v.text for v in nodo.findall('{*}v')]
        if valores:
            rv_a_rel[n] = int(valores[0])            # primer campo = _rvRel:LocalImageIdentifier en imágenes
    # 3) índice -> rId -> media
    rel_ids = []
    if 'xl/richData/richValueRel.xml' in nombres:
        for nodo in ET.fromstring(z.read('xl/richData/richValueRel.xml')):
            rel_ids.append(nodo.get(f'{{{NS["r"]}}}id'))
    rels = _rels(z, 'xl/richData/richValueRel.xml')
    # 4) celdas con vm en la hoja
    salida = []
    hoja = ET.fromstring(z.read(parte_hoja))
    for c in hoja.iter(f'{{{ns_m}}}c'):
        vm = c.get('vm')
        if not vm:
            continue
        idx_rv = vm_a_rv.get(int(vm))
        idx_rel = rv_a_rel.get(idx_rv) if idx_rv is not None else None
        if idx_rel is None or idx_rel >= len(rel_ids) or rel_ids[idx_rel] not in rels:
            salida.append({'tipo': 'en_celda', 'celda': c.get('r'), 'resuelta': False})
            continue
        salida.append({'tipo': 'en_celda', 'celda': c.get('r'), 'resuelta': True,
                       **_info_imagen(z, rels[rel_ids[idx_rel]][1])})
    return salida


# ----------------------------------------------------------------------------------------
# Hojas y celdas
# ----------------------------------------------------------------------------------------
def _valor_json(v):
    if hasattr(v, 'isoformat'):
        return v.isoformat()
    if isinstance(v, (int, float, str, bool)) or v is None:
        return v
    return str(v)


def leer_hoja(ws_f, ws_v, imagenes: list[dict]) -> dict:
    combinadas = {}
    for rango in ws_f.merged_cells.ranges:
        for fila in ws_f.iter_rows(min_row=rango.min_row, max_row=rango.max_row,
                                   min_col=rango.min_col, max_col=rango.max_col):
            for c in fila:
                combinadas[c.coordinate] = str(rango)
    img_por_celda = {}
    for im in imagenes:
        celda = im.get('celda') or im.get('celda_desde')
        if celda:
            img_por_celda.setdefault(celda, []).append(im.get('sha256') or im.get('nombre'))
    celdas = OrderedDict()
    for fila in ws_f.iter_rows():
        for c in fila:
            coord = c.coordinate
            extra = coord in combinadas or coord in img_por_celda or c.comment or c.hyperlink
            if c.value is None and not extra:
                continue
            entrada = {}
            es_formula = c.data_type == 'f'
            if es_formula:
                entrada['formula'] = str(c.value)
                entrada['valor'] = _valor_json(ws_v[coord].value)          # caché de Excel; None si nunca se calculó
            else:
                entrada['valor'] = _valor_json(c.value)
            if c.number_format and c.number_format != 'General':
                entrada['formato'] = c.number_format
            if coord in combinadas:
                entrada['combinada_en'] = combinadas[coord]
            if coord in img_por_celda:
                entrada['imagenes'] = img_por_celda[coord]
            if c.comment:
                entrada['comentario'] = c.comment.text
            if c.hyperlink and c.hyperlink.target:
                entrada['enlace'] = c.hyperlink.target
            if isinstance(entrada.get('valor'), str) and PATRON_ORDEN.search(entrada['valor']):
                entrada['sospecha_instruccion'] = True
            celdas[coord] = entrada
    return {
        'estado': ws_f.sheet_state,
        'dimension': ws_f.dimensions,
        'filas': ws_f.max_row, 'columnas': ws_f.max_column,
        'congelado': ws_f.freeze_panes,
        'filtro': ws_f.auto_filter.ref,
        'combinadas': [str(r) for r in ws_f.merged_cells.ranges],
        'filas_ocultas': sorted(k for k, d in ws_f.row_dimensions.items() if d.hidden),
        'columnas_ocultas': sorted(k for k, d in ws_f.column_dimensions.items() if d.hidden),
        'validaciones': [str(v.sqref) for v in (ws_f.data_validations.dataValidation if ws_f.data_validations else [])],
        'imagenes': imagenes,
        'celdas': celdas,
    }


def extraer(ruta: Path) -> dict:
    avisos = revisar_zip(ruta)
    wb_f = openpyxl.load_workbook(ruta, data_only=False)
    wb_v = openpyxl.load_workbook(ruta, data_only=True)
    hojas = {}
    with zipfile.ZipFile(ruta) as z:
        partes = _hoja_a_partes(z)
        for ws in wb_f.worksheets:
            imgs = imagenes_flotantes(z, partes[ws.title]) + imagenes_en_celda(z, partes[ws.title])
            hojas[ws.title] = leer_hoja(ws, wb_v[ws.title], imgs)
    dn = wb_f.defined_names
    # openpyxl >= 3.1: dict; 3.0.x: lista en .definedName
    pares = dn.items() if hasattr(dn, 'items') else [(d.name, d) for d in dn.definedName]
    nombres_definidos = {n: d.attr_text for n, d in pares if not n.startswith('_xlnm')}
    return {'esquema': ESQUEMA, 'archivo': ruta.name,
            'sha256': hashlib.sha256(ruta.read_bytes()).hexdigest(),
            'bytes': ruta.stat().st_size, 'avisos_contenedor': avisos,
            'nombres_definidos': nombres_definidos, 'hojas': hojas}


# ----------------------------------------------------------------------------------------
# Resumen medido + riesgos
# ----------------------------------------------------------------------------------------
def resumir(inter: dict) -> dict:
    hojas = []
    riesgos = list(inter['avisos_contenedor'])
    for nombre, h in inter['hojas'].items():
        formulas = [k for k, c in h['celdas'].items() if 'formula' in c]
        sin_cache = [k for k in formulas if h['celdas'][k]['valor'] is None]
        tipos = Counter(i['tipo'] for i in h['imagenes'])
        hojas.append({'hoja': nombre, 'estado': h['estado'], 'dimension': h['dimension'],
                      'celdas_con_dato': len(h['celdas']), 'combinadas': len(h['combinadas']),
                      'formulas': len(formulas), 'formulas_sin_valor_cache': len(sin_cache),
                      'imagenes': dict(tipos), 'filas_ocultas': len(h['filas_ocultas']),
                      'columnas_ocultas': len(h['columnas_ocultas']),
                      'comentarios': sum(1 for c in h['celdas'].values() if 'comentario' in c),
                      'sospecha_instruccion': sum(1 for c in h['celdas'].values() if c.get('sospecha_instruccion'))})
        if h['estado'] != 'visible':
            riesgos.append(f'hoja «{nombre}» {h["estado"]}')
        if sin_cache:
            riesgos.append(f'hoja «{nombre}»: {len(sin_cache)} fórmulas sin valor calculado (abrir y guardar en Excel)')
        if any(not i.get('resuelta', True) for i in h['imagenes']):
            riesgos.append(f'hoja «{nombre}»: imagen en celda no resuelta')
    return {'archivo': inter['archivo'], 'sha256': inter['sha256'], 'hojas': hojas, 'riesgos': riesgos}


# ----------------------------------------------------------------------------------------
# Interpretación determinista de la tabla principal (sin modelo)
# ----------------------------------------------------------------------------------------
def _precio(texto):
    m = re.search(r'\$\s*([\d,]+\.\d{2})', texto or '')
    return float(m.group(1).replace(',', '')) if m else None


def precios_por_tamano(texto) -> dict:
    """«Full $149.99 · Queen $169.99 · King $—» -> {'Full':149.99,'Queen':169.99,'King':None}."""
    salida = OrderedDict()
    for parte in re.split(r'[·|;]', texto or ''):
        m = re.match(r'\s*(Twin|Full|Queen|King|Individual|Matrimonial)\b(.*)', parte.strip(), re.I)
        if m:
            salida[m.group(1).capitalize()] = _precio(m.group(2))
    return salida


def detectar_encabezado(h: dict, max_filas=10):
    """Fila con más celdas de texto corto seguidas de filas con datos: heurística simple."""
    por_fila = {}
    for coord, c in h['celdas'].items():
        m = re.match(r'([A-Z]+)(\d+)', coord)
        por_fila.setdefault(int(m.group(2)), {})[m.group(1)] = c.get('valor')
    mejor, puntaje = None, -1
    for fila in sorted(por_fila)[:max_filas]:
        vals = por_fila[fila]
        textos = [v for v in vals.values() if isinstance(v, str) and 0 < len(v) < 40]
        siguiente = por_fila.get(fila + 1, {})
        p = len(textos) + (2 if len(siguiente) >= len(textos) - 1 else 0)
        if len(textos) == len(vals) and p > puntaje:
            mejor, puntaje = fila, p
    return mejor, por_fila


def candidatos_productos(inter: dict, hoja='Productos', fotos: Path | None = None) -> dict:
    h = inter['hojas'].get(hoja)
    if not h:
        return {'hoja': hoja, 'error': 'no existe'}
    fila_enc, por_fila = detectar_encabezado(h)
    enc = por_fila[fila_enc]
    # mapeo de columnas por palabras clave (el modelo solo entraría si esto falla: confianza < 1)
    # OJO con el orden: «Precio por tamaño» también casa con «precio»; lo específico va primero.
    claves = {'codigo': r'c[oó]digo', 'nombre': r'producto|nombre', 'precios_tamano': r'tama[nñ]o',
              'precio': r'^precio\b(?!s)', 'combo': r'combo|par', 'stock': r'stock|existenc', 'obs': r'observ|nota'}
    mapa, sin_mapear = {}, []
    for col, titulo in enc.items():
        for k, pat in claves.items():
            if k not in mapa.values() and re.search(pat, str(titulo), re.I):
                mapa[col] = k
                break
        else:
            sin_mapear.append(col)
    archivos = {}
    if fotos and fotos.exists():
        for f in fotos.iterdir():
            if f.suffix.lower() in ('.png', '.jpg', '.jpeg'):
                base = re.sub(r'_\d+$', '', f.stem).upper()
                archivos.setdefault(base, []).append(f.name)
        clave = lambda n: [int(t) if t.isdigit() else t for t in re.split(r'(\d+)', n)]  # noqa: E731
        archivos = {k: sorted(v, key=clave) for k, v in archivos.items()}
    filas, discrepancias, vistos = [], [], {}
    for fila in sorted(k for k in por_fila if k > fila_enc):
        celdas = {mapa[c]: v for c, v in por_fila[fila].items() if c in mapa}
        codigo = (str(celdas.get('codigo') or '')).strip()
        precios = precios_por_tamano(celdas.get('precios_tamano')) if celdas.get('precios_tamano') else {}
        if not precios and isinstance(celdas.get('precio'), (int, float)):
            precios = {'': float(celdas['precio'])}
        ref = f'{hoja}!A{fila}'
        reg = {'origen': ref, 'codigo': codigo or None, 'nombre': celdas.get('nombre'), 'precios': precios,
               'combo_texto': celdas.get('combo'), 'obs_origen': celdas.get('obs'),
               'fotos': archivos.get(codigo.replace('/', '-').upper(), []) if fotos else None, 'alertas': []}
        if not codigo:
            reg['alertas'].append('sin_codigo')
        if not precios or all(v is None for v in precios.values()):
            reg['alertas'].append('sin_precio')
        if any(v is None for v in precios.values()) and len(precios) > 1:
            reg['alertas'].append('tamano_sin_precio')
        ordenados = [precios.get(t) for t in TAMANOS if precios.get(t) is not None]
        if ordenados != sorted(ordenados):
            reg['alertas'].append('precios_por_tamano_no_crecientes')
        if codigo and codigo != codigo.upper():
            reg['alertas'].append('codigo_con_minusculas')
        if codigo in vistos:
            previo = vistos[codigo]
            if previo['precios'] != precios:
                reg['alertas'].append(f'duplicado_con_otro_precio_vs_{previo["origen"]}')
                discrepancias.append({'codigo': codigo, 'filas': [previo['origen'], ref],
                                      'precios': [previo['precios'], precios]})
            else:
                reg['alertas'].append(f'duplicado_identico_de_{previo["origen"]}')
        elif codigo:
            vistos[codigo] = reg
        if fotos is not None and codigo and not reg['fotos']:
            reg['alertas'].append('sin_foto_en_carpeta')
        filas.append(reg)
    conteo = Counter(a.split('_vs_')[0].split('_de_')[0] for f in filas for a in f['alertas'])
    return {'hoja': hoja, 'fila_encabezado': fila_enc, 'mapa_columnas': mapa, 'columnas_sin_mapear': sin_mapear,
            'confianza_mapa': 1.0 if not sin_mapear else 0.5, 'filas': len(filas),
            'codigos_unicos': len(vistos), 'alertas_por_tipo': dict(conteo),
            'discrepancias': discrepancias, 'registros': filas}


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument('xlsx', type=Path)
    ap.add_argument('--salida', type=Path, default=Path('salida'))
    ap.add_argument('--fotos', type=Path, default=None, help='carpeta con CODIGO.png para cruzar fotos')
    ap.add_argument('--hoja-principal', default='Productos')
    a = ap.parse_args(argv)
    inter = extraer(a.xlsx)
    resumen = resumir(inter)
    cand = candidatos_productos(inter, a.hoja_principal, a.fotos)
    a.salida.mkdir(parents=True, exist_ok=True)
    base = a.xlsx.stem
    for sufijo, datos in (('intermedio', inter), ('resumen', resumen), ('candidatos', cand)):
        (a.salida / f'{base}.{sufijo}.json').write_text(json.dumps(datos, ensure_ascii=False, indent=1) + '\n', 'utf-8')
    print(json.dumps(resumen, ensure_ascii=False, indent=1))
    print('\nCandidatos:', json.dumps({k: v for k, v in cand.items() if k != 'registros'}, ensure_ascii=False, indent=1)[:1800])
    return 0


if __name__ == '__main__':
    sys.exit(main())
