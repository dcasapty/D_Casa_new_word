"""Línea base Container: openpyxl read_only (streaming, sin imágenes) y completo (con imágenes)."""
import json, resource, sys, time
from openpyxl import load_workbook
modo, archivo = sys.argv[1], sys.argv[2]
base = resource.getrusage(resource.RUSAGE_SELF).ru_maxrss
t0 = time.perf_counter(); c0 = time.process_time()
wb = load_workbook(archivo, read_only=(modo == 'ro'), data_only=False)
filas = 0; imgs = 0
for ws in wb.worksheets:
    for _ in ws.iter_rows(values_only=True):
        filas += 1
    imgs += len(getattr(ws, '_images', []))
ms = (time.perf_counter() - t0) * 1000; cpu = (time.process_time() - c0) * 1000
fin = resource.getrusage(resource.RUSAGE_SELF).ru_maxrss
print(json.dumps({'lector': f'openpyxl-{modo}', 'archivo': archivo.split('/')[-1], 'ok': True, 'ms': round(ms), 'cpu_ms': round(cpu),
                  'rss_delta_mb': round((fin - base) / 1024, 1), 'heap_pico_mb': None, 'detalle': f'{filas} filas, {imgs} img'}))
